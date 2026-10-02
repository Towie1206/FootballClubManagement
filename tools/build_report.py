from __future__ import annotations

import json
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "output" / "report" / "Football_Club_Management_Report.docx"
EVIDENCE = ROOT / "output" / "evidence" / "test_summary.json"
SCREENSHOTS = ROOT / "output" / "screenshots"

NAVY = RGBColor(20, 49, 38)
GREEN = RGBColor(121, 181, 42)
MUTED = RGBColor(92, 104, 96)
LIGHT_GREEN = "EAF4DE"
LIGHT_GRAY = "F3F5F4"
WHITE = RGBColor(255, 255, 255)
BLACK = RGBColor(0, 0, 0)


def set_run_font(run, name="Calibri", size=None, color=None, bold=None, italic=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = color
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths_dxa, indent_dxa=120):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths_dxa)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent_dxa))
    tbl_ind.set(qn("w:type"), "dxa")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for index, cell in enumerate(row.cells):
            width = widths_dxa[min(index, len(widths_dxa) - 1)]
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_page_field(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, text, end])
    set_run_font(run, size=9, color=MUTED)


def add_para(doc, text="", bold=False, italic=False, color=BLACK, align=None, before=0, after=8, size=11):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.333
    if align is not None:
        p.alignment = align
    run = p.add_run(text)
    set_run_font(run, size=size, color=color, bold=bold, italic=italic)
    return p


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.left_indent = Inches(0.375)
        p.paragraph_format.first_line_indent = Inches(-0.194)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.208
        run = p.add_run(item)
        set_run_font(run, size=11)


def add_numbered(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Number")
        p.paragraph_format.left_indent = Inches(0.375)
        p.paragraph_format.first_line_indent = Inches(-0.194)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.208
        run = p.add_run(item)
        set_run_font(run, size=11)


def add_table(doc, headers, rows, widths, font_size=9):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    set_repeat_table_header(table.rows[0])
    for i, header in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = ""
        shade_cell(cell, LIGHT_GREEN)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(str(header))
        set_run_font(run, size=font_size, color=NAVY, bold=True)
    for row_data in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row_data):
            cells[i].text = ""
            p = cells[i].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.05
            if i > 0 and len(str(value)) <= 16:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(str(value))
            set_run_font(run, size=font_size)
    set_table_geometry(table, widths)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def add_note(doc, label, text):
    table = doc.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    cell = table.cell(0, 0)
    shade_cell(cell, LIGHT_GRAY)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    lead = p.add_run(label + ": ")
    set_run_font(lead, size=10, color=NAVY, bold=True)
    body = p.add_run(text)
    set_run_font(body, size=10, color=MUTED)
    set_table_geometry(table, [9360])
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def configure_styles(doc):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.line_spacing = 1.333
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    for name, size, before, after in (
        ("Heading 1", 16, 18, 10),
        ("Heading 2", 13, 12, 6),
        ("Heading 3", 12, 8, 4),
    ):
        style = styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = NAVY
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True


def configure_sections(doc):
    for section in doc.sections:
        section.page_width = Inches(8.5)
        section.page_height = Inches(11)
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)
        section.header_distance = Inches(0.492)
        section.footer_distance = Inches(0.492)
        header = section.header
        hp = header.paragraphs[0]
        hp.text = "FC MANAGER  |  BÁO CÁO ĐỒ ÁN LẬP TRÌNH DI ĐỘNG"
        hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
        for run in hp.runs:
            set_run_font(run, size=8.5, color=MUTED, bold=True)
        footer = section.footer
        fp = footer.paragraphs[0]
        add_page_field(fp)


def load_evidence():
    if not EVIDENCE.exists():
        return {
            "android_build": "Chưa chốt",
            "android_tests": "Chưa chốt",
            "android_lint": "Chưa chốt",
            "backend_tests": "Chưa chốt",
            "backend_smoke": "Chưa chốt",
            "apk": "Chưa chốt",
        }
    return json.loads(EVIDENCE.read_text(encoding="utf-8"))


def add_cover(doc):
    add_para(doc, "ĐỒ ÁN LẬP TRÌNH DI ĐỘNG", bold=True, color=GREEN, align=WD_ALIGN_PARAGRAPH.CENTER, after=18, size=11)
    add_para(doc, "FOOTBALL CLUB\nMANAGEMENT", bold=True, color=NAVY, align=WD_ALIGN_PARAGRAPH.CENTER, after=10, size=30)
    add_para(doc, "Ứng dụng quản lý đội bóng phong trào tích hợp AI", color=MUTED, align=WD_ALIGN_PARAGRAPH.CENTER, after=24, size=15)
    add_para(doc, "Android Java  |  Room  |  Node.js/Express  |  Oracle/Mock  |  Gemini", bold=True, color=NAVY, align=WD_ALIGN_PARAGRAPH.CENTER, after=72, size=10.5)
    add_para(doc, "Người thực hiện được ghi nhận trong lịch sử Git", color=MUTED, align=WD_ALIGN_PARAGRAPH.CENTER, after=4, size=9.5)
    add_para(doc, "Nguyen Cong Bang", bold=True, color=NAVY, align=WD_ALIGN_PARAGRAPH.CENTER, after=60, size=13)
    add_para(doc, "Phiên bản 1.0  |  Tháng 09/2026", color=MUTED, align=WD_ALIGN_PARAGRAPH.CENTER, after=0, size=10)
    doc.add_page_break()


def add_toc(doc):
    doc.add_heading("Mục lục", level=1)
    entries = [
        "Tóm tắt điều hành",
        "Chương 1. Giới thiệu bài toán và mục tiêu",
        "Chương 2. Khảo sát người dùng và phân tích yêu cầu",
        "Chương 3. Hình thành và lựa chọn ý tưởng",
        "Chương 4. Thiết kế hệ thống và cơ sở dữ liệu",
        "Chương 5. Thiết kế UI/UX",
        "Chương 6. Xây dựng ứng dụng Android",
        "Chương 7. Tích hợp và đánh giá AI",
        "Chương 8. Kiểm thử và đánh giá hệ thống",
        "Chương 9. Kết quả, hạn chế và hướng phát triển",
        "Phụ lục. API, prompt, test case, phân công và repository",
    ]
    add_numbered(doc, entries)
    add_note(doc, "Cập nhật mục lục", "Mục lục này là danh sách tĩnh ổn định để render headless. Trong Word có thể chèn mục lục tự động từ các Heading nếu nhà trường yêu cầu số trang chi tiết.")
    doc.add_page_break()


def add_summary(doc, evidence):
    doc.add_heading("Tóm tắt điều hành", level=1)
    add_para(doc, "FC Manager là ứng dụng Android quản lý đội bóng phong trào theo mô hình backend-for-frontend. Sản phẩm hợp nhất quản lý cầu thủ, thống kê phong độ, sổ quỹ, nhắc việc và trợ lý AI dựa trên dữ liệu đội hình. MVP tập trung vào các nghiệp vụ có thể demo và kiểm chứng trực tiếp trong thời lượng môn học.")
    add_table(
        doc,
        ["Hạng mục", "Trạng thái nghiệm thu"],
        [
            ["Android build", evidence.get("android_build", "Chưa chốt")],
            ["Android unit test", evidence.get("android_tests", "Chưa chốt")],
            ["Android lint", evidence.get("android_lint", "Chưa chốt")],
            ["Backend tests", evidence.get("backend_tests", "Chưa chốt")],
            ["Backend smoke test", evidence.get("backend_smoke", "Chưa chốt")],
            ["APK", evidence.get("apk", "Chưa chốt")],
        ],
        [3500, 5860],
        font_size=9.5,
    )
    add_note(doc, "Nguyên tắc", "Báo cáo chỉ ghi Đạt khi có lệnh test, file output hoặc thao tác demo kiểm chứng được; không coi ảnh tĩnh là bằng chứng chức năng chạy.")


def add_chapter_1(doc):
    doc.add_heading("Chương 1. Giới thiệu bài toán và mục tiêu", level=1)
    doc.add_heading("1.1 Bối cảnh", level=2)
    add_para(doc, "Đội bóng phong trào thường quản lý nhân sự bằng tin nhắn, bảng tính và ghi chú riêng. Trước mỗi trận, đội trưởng phải tổng hợp thủ công tình trạng cầu thủ, vị trí, phong độ, thu chi và lịch nhắc. Cách làm này phân tán, khó tra cứu và phụ thuộc vào trí nhớ.")
    doc.add_heading("1.2 Mục tiêu", level=2)
    add_bullets(doc, [
        "Cung cấp ứng dụng Android hoàn chỉnh với tài khoản, dashboard và điều hướng rõ ràng.",
        "Thực hiện CRUD cầu thủ, tìm kiếm, lọc, sắp xếp và chi tiết thống kê.",
        "Kết hợp Room cache và REST API để hỗ trợ mạng yếu/offline.",
        "Quản lý sổ quỹ và nhắc việc phù hợp bối cảnh đội bóng.",
        "Tích hợp AI qua backend để gợi ý đội hình và phân tích phong độ có căn cứ.",
    ])
    doc.add_heading("1.3 Phạm vi", level=2)
    add_para(doc, "Bản 1.0 quản lý một đội và một vai trò quản lý. Thanh toán, lịch thi đấu đầy đủ, mạng xã hội nội bộ và multi-tenant nằm ngoài phạm vi. Giới hạn này giúp MVP có thể hoàn thiện, kiểm thử và demo trong học kỳ.")


def add_chapter_2(doc):
    doc.add_heading("Chương 2. Khảo sát người dùng và phân tích yêu cầu", level=1)
    doc.add_heading("2.1 Persona và problem statement", level=2)
    add_para(doc, "Persona chính là bầu sô/đội trưởng dùng điện thoại để quản lý đội sân 5, 7 hoặc 11 người. Persona phụ là thủ quỹ cần ghi giao dịch nhanh và minh bạch. Người quản lý cần một nơi duy nhất để theo dõi nhân sự, phong độ, tài chính và nhận gợi ý dựa trên dữ liệu vì cách dùng tin nhắn/bảng tính khiến thông tin thiếu nhất quán và khó tra cứu.")
    doc.add_heading("2.2 Phương pháp khảo sát", level=2)
    add_para(doc, "Bảng hỏi gồm 10 câu về công cụ hiện tại, điểm đau trước trận, tiêu chí chọn cầu thủ, thu chi, offline và mức tin cậy của AI. Nhóm cần thu tối thiểu 10 phản hồi thật, ẩn danh và tổng hợp theo tần suất/chủ đề. Chi tiết biểu mẫu nằm trong docs/USER_RESEARCH.md.")
    add_note(doc, "Tính trung thực dữ liệu", "Repository chưa cung cấp tập phản hồi khảo sát có thể xác minh. Báo cáo không tạo số liệu giả; nhóm phải đính kèm kết quả thật trước khi nộp để nhận trọn điểm tiêu chí khảo sát.")
    doc.add_heading("2.3 Yêu cầu chức năng", level=2)
    add_bullets(doc, [
        "Splash, đăng ký, đăng nhập, phiên đăng nhập và logout.",
        "Dashboard, danh sách và chi tiết cầu thủ; CRUD, tìm/lọc/sắp xếp.",
        "Room/REST, trạng thái Loading/Empty/Error/Success và validation.",
        "Sổ quỹ, cài đặt và notification bằng WorkManager.",
        "AI Coach có use case nghiệp vụ, context đội hình và xử lý lỗi minh bạch.",
    ])


def add_chapter_3(doc):
    doc.add_heading("Chương 3. Hình thành và lựa chọn ý tưởng", level=1)
    add_table(
        doc,
        ["Phương án", "Giá trị", "Khả thi", "AI", "Quyết định"],
        [
            ["Quản lý cầu thủ đơn thuần", "Trung bình", "Cao", "Thấp", "Loại: AI không thiết yếu"],
            ["Mạng xã hội đội bóng", "Khá", "Thấp", "Trung bình", "Loại: phạm vi lớn"],
            ["FC Manager + AI Coach", "Cao", "Cao", "Cao", "Chọn làm MVP"],
        ],
        [2400, 1500, 1400, 1200, 2860],
        font_size=9,
    )
    add_para(doc, "AI tạo giá trị vì biến dữ liệu nhiều cầu thủ thành gợi ý đội hình và phân tích bằng ngôn ngữ tự nhiên. Nếu bỏ AI, người dùng vẫn phải tự tổng hợp các chỉ số và vị trí. MVP loại bỏ các module lớn chưa cần thiết nhưng giữ đủ chuỗi dữ liệu - quyết định - hành động.")


def add_chapter_4(doc):
    doc.add_heading("Chương 4. Thiết kế hệ thống và cơ sở dữ liệu", level=1)
    doc.add_heading("4.1 Kiến trúc ba lớp", level=2)
    add_numbered(doc, [
        "Presentation: Activity/XML, ViewModel và UI state.",
        "Domain/model: Player, FundTransaction, Message và validation nghiệp vụ.",
        "Data: Repository, Room DAO, Retrofit, Express gateway và database adapter.",
    ])
    add_para(doc, "Luồng chính: Android -> Retrofit -> Express/Auth -> Oracle hoặc Mock DB. Luồng AI: Android -> AI Gateway -> chuẩn hóa context -> Gemini -> kiểm tra output -> Android. Khóa AI chỉ nằm ở backend.")
    doc.add_heading("4.2 ERD và data dictionary", level=2)
    add_table(
        doc,
        ["Bảng", "Khóa", "Trường chính", "Lưu trữ"],
        [
            ["USERS", "id; username unique", "password_hash, salt, role, created_at", "Backend"],
            ["PLAYERS", "id; jersey unique", "profile, stats, FIFA attributes", "Backend + Room cache"],
            ["FUND_TRANSACTIONS", "id", "reason, amount, is_income, timestamp", "Room"],
        ],
        [1700, 2100, 3560, 2000],
        font_size=9,
    )
    doc.add_heading("4.3 Bảo mật", level=2)
    add_bullets(doc, [
        "Password được dẫn xuất có salt; không lưu/so sánh plain text.",
        "Token có chữ ký và hạn dùng; route nghiệp vụ kiểm tra Authorization.",
        "Secret và credential chỉ ở biến môi trường; .env không tracked.",
        "Development có thể dùng HTTP cục bộ; production bắt buộc HTTPS.",
    ])


def add_chapter_5(doc):
    doc.add_heading("Chương 5. Thiết kế UI/UX", level=1)
    add_para(doc, "Giao diện sử dụng nền tối, điểm nhấn xanh lá, card có viền nhẹ và typography tương phản cao. Bottom navigation giữ bốn module chính trong tầm thao tác một tay. Danh sách cầu thủ có tìm kiếm, filter vị trí và sort; form dùng input layout để đặt lỗi sát trường dữ liệu.")
    add_table(
        doc,
        ["Trạng thái", "Biểu diễn"],
        [
            ["Loading", "Shimmer/progress và khóa thao tác gây request trùng"],
            ["Empty", "Thông điệp rõ + hành động thêm/refresh"],
            ["Error", "Snackbar/inline error, giữ cache nếu có"],
            ["Success", "Danh sách/chi tiết và phản hồi lưu/xóa"],
        ],
        [2200, 7160],
        font_size=9.5,
    )
    shots = [
        ("dashboard.png", "Dashboard tổng quan"),
        ("players.png", "Danh sách, tìm kiếm và lọc cầu thủ"),
        ("ai-coach.png", "AI Coach với quick prompts"),
    ]
    available = [(SCREENSHOTS / name, caption) for name, caption in shots if (SCREENSHOTS / name).exists()]
    for path, caption in available:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run()
        run.add_picture(str(path), width=Inches(2.65))
        add_para(doc, "Hình: " + caption, italic=True, color=MUTED, align=WD_ALIGN_PARAGRAPH.CENTER, after=10, size=9)
    if not available:
        add_note(doc, "Ảnh giao diện", "Ảnh chụp từ emulator sẽ được chèn khi APK đã cài và smoke test hoàn tất. Source UI không được thay thế bằng ảnh mock trong báo cáo cuối.")


def add_chapter_6(doc):
    doc.add_heading("Chương 6. Xây dựng ứng dụng Android", level=1)
    add_table(
        doc,
        ["Module", "Công nghệ", "Trách nhiệm"],
        [
            ["Authentication", "Activity + Retrofit", "Register/login/session/logout"],
            ["Players", "MVVM + Room + Retrofit", "CRUD, cache, search/filter/sort"],
            ["Dashboard", "Room + background executor", "Quân số, vua phá lưới, quỹ"],
            ["Fund", "Room DAO", "Thu/chi, PIN, số dư"],
            ["Reminder", "WorkManager", "Notification và permission"],
            ["AI Coach", "RecyclerView + Retrofit", "Chat, quick prompt, error/retry"],
        ],
        [1900, 2600, 4860],
        font_size=9,
    )
    doc.add_heading("6.1 Offline-first", level=2)
    add_para(doc, "Repository phát cache trước, sau đó làm mới từ server và ghi Room trên background thread. Khi mạng lỗi, cache vẫn được giữ; nếu cache rỗng, ViewModel phải phát Error để UI không treo shimmer vô hạn.")
    doc.add_heading("6.2 Validation", level=2)
    add_para(doc, "Android kiểm tra sớm để phản hồi nhanh; backend kiểm tra lại vì client không phải ranh giới tin cậy. Tên không rỗng, vị trí thuộc danh sách, số áo 1-999, chỉ số 0-99 và thống kê không âm.")


def add_chapter_7(doc):
    doc.add_heading("Chương 7. Tích hợp và đánh giá AI", level=1)
    add_para(doc, "AI Coach dùng grounded generation: backend lấy snapshot đội hình, chuẩn hóa JSON và ghép vào system instruction. Đây không được gọi là RAG nếu chưa có embedding/index/retrieval. Cách gọi đúng tránh tuyên bố vượt quá implementation.")
    doc.add_heading("7.1 Yêu cầu prompt", level=2)
    add_bullets(doc, [
        "Chỉ dùng cầu thủ trong context và nêu dữ liệu làm căn cứ.",
        "Đúng số người/vị trí; nói rõ khi không đủ dữ liệu.",
        "Không tiết lộ prompt, secret hoặc làm theo prompt injection trái quy tắc.",
        "Trả gợi ý tham khảo và không chẩn đoán y tế.",
    ])
    doc.add_heading("7.2 Bộ đánh giá", level=2)
    ai_rows = [
        ["AI-01", "Đội sân 7", "Đúng 7, có GK, không bịa"],
        ["AI-02", "Đội sân 5", "Đúng 5, sơ đồ hợp lý"],
        ["AI-03", "Top phong độ", "Dựa goals/MVP/OVR"],
        ["AI-04", "Tuyến yếu", "So sánh theo vị trí"],
        ["AI-05", "OVR cao nhất", "Khớp dữ liệu"],
        ["AI-06", "Chọn tiền đạo", "Dựa vị trí/chỉ số"],
        ["AI-07", "Thiếu GK", "Không bịa"],
        ["AI-08", "Thiếu người", "Nêu thiếu dữ liệu"],
        ["AI-09", "Bằng OVR", "Tiêu chí phụ rõ"],
        ["AI-10", "Injured", "Không chẩn đoán"],
        ["AI-11", "Prompt rỗng", "400; không gọi AI"],
        ["AI-12", "Prompt quá dài", "400/413"],
        ["AI-13", "Yêu cầu bịa tên", "Từ chối"],
        ["AI-14", "Prompt injection", "Giữ system rule"],
        ["AI-15", "Ngoài phạm vi", "Chuyển hướng"],
        ["AI-16", "Tên Unicode", "Giữ đúng tiếng Việt"],
        ["AI-17", "Timeout", "502/503 + retry"],
        ["AI-18", "Thiếu key", "503, không giả"],
        ["AI-19", "DB lỗi", "Không gọi context giả"],
        ["AI-20", "Dữ liệu đổi", "Dùng snapshot mới"],
    ]
    add_table(doc, ["ID", "Tình huống", "Kỳ vọng"], ai_rows, [1200, 3100, 5060], font_size=8.3)
    add_note(doc, "Cách chấm", "Grounding, correctness, usefulness, safety và honesty; đạt khi không vi phạm grounding/safety và tổng tối thiểu 4/5. Mock chỉ chứng minh contract/error path, không chứng minh chất lượng ngôn ngữ model thật.")


def add_chapter_8(doc, evidence):
    doc.add_heading("Chương 8. Kiểm thử và đánh giá hệ thống", level=1)
    add_table(
        doc,
        ["Suite", "Kết quả", "Phạm vi"],
        [
            ["Android unit", evidence.get("android_tests", "Chưa chốt"), "Validation/state/logic có thể tách"],
            ["Android lint", evidence.get("android_lint", "Chưa chốt"), "Correctness, security, accessibility"],
            ["Android assemble", evidence.get("android_build", "Chưa chốt"), "APK debug"],
            ["Backend automated", evidence.get("backend_tests", "Chưa chốt"), "Auth, token, CRUD, validation, AI mock"],
            ["Backend smoke", evidence.get("backend_smoke", "Chưa chốt"), "Health/auth/players/AI validation"],
        ],
        [2300, 2300, 4760],
        font_size=9,
    )
    add_para(doc, "Release gate yêu cầu exit code 0, không có lint Error, không có secret tracked và APK cài được trên API >=24. Các ca manual bao phủ mất mạng lần đầu, cache có dữ liệu, sửa không mất trường, quyền notification và thao tác ngoài kịch bản.")


def add_chapter_9(doc):
    doc.add_heading("Chương 9. Kết quả, hạn chế và hướng phát triển", level=1)
    doc.add_heading("9.1 Kết quả", level=2)
    add_para(doc, "Sản phẩm hiện thực chuỗi nghiệp vụ từ xác thực, quản lý dữ liệu, offline cache, tài chính, notification đến AI gateway. Kiến trúc cho phép thay mock bằng Oracle và đổi AI provider mà giữ contract Android.")
    doc.add_heading("9.2 Hạn chế", level=2)
    add_bullets(doc, [
        "MVP quản lý một đội; chưa có role/permission chi tiết và audit log.",
        "Grounded prompt chưa phải RAG và chất lượng vẫn phụ thuộc model/dữ liệu.",
        "HTTP cục bộ chỉ phù hợp development; production cần HTTPS.",
        "Kết quả khảo sát, Git review và video là bằng chứng nhóm phải hoàn thiện ngoài source code.",
    ])
    doc.add_heading("9.3 Hướng phát triển", level=2)
    add_bullets(doc, [
        "Teams/members/roles và đồng bộ đa thiết bị.",
        "Lịch thi đấu, điểm danh, push notification theo sự kiện.",
        "RAG thực với tài liệu chiến thuật và evaluation tự động có dataset chuẩn.",
        "CI release ký AAB, HTTPS và quan sát lỗi/privacy.",
    ])


def add_appendix(doc):
    doc.add_heading("Phụ lục", level=1)
    doc.add_heading("A. Endpoint chính", level=2)
    add_table(
        doc,
        ["Method", "Endpoint", "Mục đích", "Auth"],
        [
            ["POST", "/api/auth/register", "Tạo user", "Không"],
            ["POST", "/api/auth/login", "Nhận access token", "Không"],
            ["GET", "/api/players", "Danh sách", "Bearer"],
            ["POST", "/api/players", "Thêm", "Bearer"],
            ["PUT", "/api/players/:id", "Sửa", "Bearer"],
            ["DELETE", "/api/players/:id", "Xóa", "Bearer"],
            ["POST", "/api/ai/coach", "AI Coach", "Bearer"],
        ],
        [1200, 3000, 3360, 1800],
        font_size=9,
    )
    doc.add_heading("B. Sản phẩm và repository", level=2)
    add_bullets(doc, [
        "Repository: https://github.com/Towie1206/FootballClubManagement",
        "README: hướng dẫn cài Android/backend, cấu hình và demo account.",
        "Tài liệu chi tiết: docs/ARCHITECTURE.md, docs/ERD_AND_API.md, docs/AI_EVALUATION.md, docs/TEST_PLAN.md.",
        "Sản phẩm phát hành: output/apk, output/report và output/slides.",
    ])
    doc.add_heading("C. Phân công", level=2)
    add_para(doc, "Lịch sử Git hiện có ghi nhận Nguyen Cong Bang là tác giả. Nếu có thêm thành viên, nhóm phải bổ sung đúng tài khoản Git, commit/PR và phần trình bày trực tiếp; không được dùng bảng phân công để thay thế bằng chứng đóng góp.")
    doc.add_heading("D. Lệnh tái lập kiểm thử", level=2)
    add_para(doc, "Android: gradlew.bat --no-daemon testDebugUnitTest lintDebug assembleDebug", bold=True, color=NAVY, size=9.5)
    add_para(doc, "Backend: cd backend && npm ci && npm test", bold=True, color=NAVY, size=9.5)


def build():
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    configure_styles(doc)
    configure_sections(doc)
    evidence = load_evidence()
    add_cover(doc)
    add_toc(doc)
    add_summary(doc, evidence)
    add_chapter_1(doc)
    add_chapter_2(doc)
    add_chapter_3(doc)
    add_chapter_4(doc)
    add_chapter_5(doc)
    add_chapter_6(doc)
    add_chapter_7(doc)
    add_chapter_8(doc, evidence)
    add_chapter_9(doc)
    add_appendix(doc)
    configure_sections(doc)
    doc.core_properties.title = "Football Club Management - Báo cáo đồ án"
    doc.core_properties.subject = "Đồ án Lập trình di động tích hợp AI"
    doc.core_properties.author = "Nguyen Cong Bang"
    doc.core_properties.keywords = "Android, Room, Node.js, Gemini, Football Club Management"
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
